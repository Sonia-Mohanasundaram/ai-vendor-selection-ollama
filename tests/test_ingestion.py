from __future__ import annotations

import unittest
from io import BytesIO

from docx import Document
import pandas as pd

from core.extraction import extract_fields
from core.ingestion import ingest_bytes
from core.normalization import normalize


class DocumentIngestionTests(unittest.TestCase):
    def test_docx_paragraphs_and_quotation_table_are_extracted(self) -> None:
        document = Document()
        document.add_paragraph("Payment Terms: Net 45 days")
        table = document.add_table(rows=1, cols=4)
        for cell, value in zip(table.rows[0].cells, ["Vendor", "Quantity", "Unit Price", "Delivery"]):
            cell.text = value
        row = table.add_row()
        for cell, value in zip(row.cells, ["Acme Supplies", "2", "INR 1,250.50", "5 days"]):
            cell.text = value

        stream = BytesIO()
        document.save(stream)
        text = ingest_bytes(stream.getvalue(), "quote.docx")
        extracted = extract_fields(text, "quote.docx")
        normalized = normalize(extracted)

        self.assertEqual(extracted["vendor_name"], "Acme Supplies")
        self.assertEqual(extracted["quantity"], 2)
        self.assertEqual(extracted["unit_price"], 1250.5)
        self.assertEqual(normalized["delivery_days"], 5)
        self.assertEqual(normalized["payment_terms_days"], 45)

    def test_xlsx_quotation_row_uses_the_same_field_mapping_as_csv(self) -> None:
        workbook = BytesIO()
        pd.DataFrame(
            [
                {
                    "vendor_name": "Acme Supplies",
                    "quantity": 2,
                    "unit_price": 1250.5,
                    "gst_percent": 18,
                    "delivery_days": 5,
                    "payment_terms_days": 30,
                    "freight_cost": 0,
                    "currency": "INR",
                }
            ]
        ).to_excel(workbook, index=False)

        text = ingest_bytes(workbook.getvalue(), "quote.xlsx")
        quote = normalize(extract_fields(text, "quote.xlsx"))
        self.assertEqual(quote["vendor_name"], "Acme Supplies")
        self.assertEqual(quote["quantity"], 2)
        self.assertEqual(quote["unit_price"], 1250.5)
        self.assertEqual(quote["total_price"], 2501.0)
        self.assertEqual(quote["delivery_days"], 5)
        self.assertEqual(quote["payment_terms_days"], 30)


if __name__ == "__main__":
    unittest.main()
